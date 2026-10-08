"""Operations for summarizing employee attendance reports."""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


EMPLOYEE_PATTERN = re.compile(r"^\s*(\d{4,}[PG]?)\s*\((.+)\)\s*$", re.IGNORECASE)


@dataclass
class AttendanceRecord:
    nik: str
    name: str
    absences: int = 0
    discipline: int = 0
    descriptions: list[str] = field(default_factory=list)

    def as_row(self) -> dict[str, Any]:
        return {
            "NIK": self.nik,
            "NAMA": self.name,
            "ABSENSI": self.absences or None,
            "SKOR ABSENSI": 4 - self.absences,
            "KEDISIPLINAN": self.discipline or None,
            "SKOR KEDISIPLINAN": 4 - self.discipline,
            "KETERANGAN": ", ".join(self.descriptions),
        }


def _read_rows(filename: str, file_bytes: bytes) -> list[tuple[Any, ...]]:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if extension == "xls":
        import xlrd
        from xlrd.biffh import XLRDError

        try:
            workbook = xlrd.open_workbook(file_contents=file_bytes)
        except XLRDError as exc:
            raise ValueError(f"File Excel .xls tidak dapat dibaca: {exc}") from exc
        try:
            sheet = workbook.sheet_by_index(0)
        except IndexError as exc:
            raise ValueError("File Excel .xls tidak memiliki sheet.") from exc
        return [tuple(sheet.row_values(row)) for row in range(sheet.nrows)]

    if extension == "xlsx":
        try:
            workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except (OSError, ValueError, zipfile.BadZipFile) as exc:
            raise ValueError(f"File Excel .xlsx tidak dapat dibaca: {exc}") from exc
        try:
            if not workbook.worksheets:
                raise ValueError("File Excel .xlsx tidak memiliki sheet.")
            return list(workbook.worksheets[0].iter_rows(values_only=True))
        finally:
            workbook.close()

    raise ValueError("Format file tidak didukung. Gunakan file .xls atau .xlsx.")


def summarize_attendance(filename: str, file_bytes: bytes) -> list[AttendanceRecord]:
    """Read a leave report and aggregate classified events once per employee."""
    records: dict[str, AttendanceRecord] = {}
    names: dict[str, str] = {}

    for row in _read_rows(filename, file_bytes):
        employee_cell = str(row[0] or "").strip() if row else ""
        match = EMPLOYEE_PATTERN.match(employee_cell)
        if not match:
            continue

        nik = match.group(1).upper()
        name = match.group(2).strip()
        if not name:
            raise ValueError(f"Baris dengan NIK {nik} tidak memiliki nama karyawan.")
        description = str(row[10] or "").strip().upper() if len(row) > 10 else ""
        if not description:
            continue

        if "SAKIT SURAT DOKTER" in description:
            category = "absence"
            code = "SD"
        elif "ALPHA" in description:
            category = "absence"
            code = "ALPHA"
        elif "IJIN" in description:
            category = "discipline"
            if "TERLAMBAT" in description:
                code = "TERLAMBAT"
            elif "SETENGAH HARI" in description:
                code = "SETENGAH HARI"
            elif "MENINGGALKAN" in description:
                code = "MENINGGALKAN"
            else:
                code = re.sub(r"^IJIN\s*", "", description).strip()
        else:
            continue

        normalized_name = " ".join(name.casefold().split())
        previous_nik = names.get(normalized_name)
        if previous_nik is not None and previous_nik != nik:
            raise ValueError(
                f"Nama karyawan {name!r} ditemukan dengan NIK berbeda "
                f"({previous_nik} dan {nik}); periksa data sumber."
            )
        names[normalized_name] = nik

        record = records.get(nik)
        if record is None:
            record = AttendanceRecord(nik=nik, name=name)
            records[nik] = record
        elif " ".join(record.name.casefold().split()) != normalized_name:
            raise ValueError(
                f"NIK {nik} ditemukan dengan nama berbeda "
                f"({record.name!r} dan {name!r}); periksa data sumber."
            )

        if category == "absence":
            record.absences += 1
        else:
            record.discipline += 1
        record.descriptions.append(code)

    if not records:
        raise ValueError(
            "Tidak ditemukan baris karyawan dengan NIK dan keterangan absensi/izin "
            "yang dikenali."
        )
    return list(records.values())


def create_attendance_workbook(records: list[AttendanceRecord]) -> bytes:
    """Create the formatted attendance summary workbook."""
    headers = [
        "NIK",
        "NAMA",
        "ABSENSI",
        "SKOR",
        "KEDISIPLINAN",
        "SKOR",
        "KETERANGAN",
    ]
    workbook = Workbook()
    sheet = workbook.create_sheet("Monitoring Absensi")
    workbook.remove(workbook.worksheets[0])
    sheet.append(headers)
    for record in records:
        sheet.append(
            [
                record.nik,
                record.name,
                record.absences or None,
                4 - record.absences,
                record.discipline or None,
                4 - record.discipline,
                ", ".join(record.descriptions),
            ]
        )

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {
        "A": 14,
        "B": 28,
        "C": 13,
        "D": 16,
        "E": 18,
        "F": 22,
        "G": 38,
    }.items():
        sheet.column_dimensions[column].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
