import io
import unittest

from openpyxl import Workbook, load_workbook

from mvc.models.attendance_operations import (
    create_attendance_workbook,
    summarize_attendance,
)


def make_report(rows):
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def event(employee, description):
    row = [None] * 11
    row[0] = employee
    row[10] = description
    return row


class AttendanceOperationsTests(unittest.TestCase):
    def test_aggregates_categories_and_deduplicates_employees(self):
        report = make_report(
            [
                event("1234P(ARI)", "SAKIT SURAT DOKTER"),
                event("1234P(ARI)", "ALPHA / TANPA KETERANGAN"),
                event("56789(SITI)", "IJIN TERLAMBAT MASUK KERJA"),
                event("9012G(BUDI)", "IJIN SETENGAH HARI"),
                event("9012G(BUDI)", "IJIN MENINGGALKAN PEKERJAAN"),
                event("1122(NINA)", "CUTI TAHUNAN"),
            ]
        )

        records = summarize_attendance("report.xlsx", report)

        self.assertEqual([record.nik for record in records], ["1234P", "56789", "9012G"])
        self.assertEqual(records[0].absences, 2)
        self.assertEqual(records[0].discipline, 0)
        self.assertEqual(records[0].as_row()["SKOR ABSENSI"], 2)
        self.assertIsNone(records[0].as_row()["KEDISIPLINAN"])
        self.assertEqual(records[0].descriptions, ["SD", "ALPHA"])
        self.assertEqual(records[1].descriptions, ["TERLAMBAT"])
        self.assertEqual(records[2].descriptions, ["SETENGAH HARI", "MENINGGALKAN"])

    def test_rejects_duplicate_name_with_different_nik(self):
        report = make_report(
            [
                event("1234(ARI)", "ALPHA"),
                event("5678(ARI)", "IJIN TERLAMBAT"),
            ]
        )
        with self.assertRaisesRegex(ValueError, "NIK berbeda"):
            summarize_attendance("report.xlsx", report)

    def test_exports_requested_columns_and_scores(self):
        report = make_report(
            [
                event("1234(ARI)", "SAKIT SURAT DOKTER"),
                event("1234(ARI)", "IJIN TERLAMBAT"),
            ]
        )
        records = summarize_attendance("report.xlsx", report)
        result = create_attendance_workbook(records)
        workbook = load_workbook(io.BytesIO(result), read_only=True, data_only=True)
        try:
            rows = list(workbook.worksheets[0].iter_rows(values_only=True))
        finally:
            workbook.close()

        self.assertEqual(
            rows[0],
            ("NIK", "NAMA", "ABSENSI", "SKOR", "KEDISIPLINAN", "SKOR", "KETERANGAN"),
        )
        self.assertEqual(rows[1], ("1234", "ARI", 1, 3, 1, 3, "SD, TERLAMBAT"))

    def test_rejects_unsupported_file_type(self):
        with self.assertRaisesRegex(ValueError, "Format file tidak didukung"):
            summarize_attendance("report.csv", b"")


if __name__ == "__main__":
    unittest.main()
